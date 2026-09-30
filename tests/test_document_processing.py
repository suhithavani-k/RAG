from io import BytesIO

import pymupdf
import pytest
from docx import Document

from core.chunker import create_chunks
from core.document_loader import DocumentLoadError, load_document
from core.document_loader import DocumentPage
from core.text_cleaner import clean_text


def test_clean_text_preserves_meaning_and_joins_hyphenated_lines():
    assert clean_text("  seman-\ntic   search\x00\n\n\nfor documents. ") == "semantic search\n\nfor documents."


def test_pdf_extraction_keeps_page_numbers():
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Eligibility requires a completed application.")
    content = pdf.tobytes()
    pdf.close()

    extracted = load_document("syllabus.pdf", content)
    assert extracted[0].page_number == 1
    assert "completed application" in extracted[0].text


def test_docx_extraction_includes_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("Program requirements")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Minimum GPA"
    table.cell(0, 1).text = "3.0"
    stream = BytesIO()
    document.save(stream)

    extracted = load_document("requirements.docx", stream.getvalue())
    assert "Program requirements" in extracted[0].text
    assert "Minimum GPA | 3.0" in extracted[0].text
    assert extracted[0].page_number is None


def test_empty_and_unsupported_documents_fail_clearly():
    empty_pdf = pymupdf.open()
    empty_pdf.new_page()
    blank_pdf_bytes = empty_pdf.tobytes()
    empty_pdf.close()
    with pytest.raises(DocumentLoadError, match="No selectable text"):
        load_document("empty.pdf", blank_pdf_bytes)
    with pytest.raises(DocumentLoadError, match="Unsupported"):
        load_document("notes.txt", b"some text")


def test_chunks_include_source_page_and_validated_overlap():
    pages = [DocumentPage(" ".join(f"word{i}" for i in range(55)), page_number=2)]
    chunks = create_chunks(pages, "syllabus.pdf", chunk_size=25, overlap=5)
    assert len(chunks) >= 2
    assert chunks[0].source == "syllabus.pdf"
    assert chunks[0].page_number == 2
    assert chunks[0].chunk_id.startswith("syllabus-2-")
    with pytest.raises(ValueError):
        create_chunks(pages, "syllabus.pdf", chunk_size=25, overlap=25)
