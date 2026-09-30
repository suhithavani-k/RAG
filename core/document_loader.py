from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document

from core.text_cleaner import clean_text


@dataclass(frozen=True)
class DocumentPage:
    text: str
    page_number: int | None = None


class DocumentLoadError(ValueError):
    """Raised when a supported document cannot provide indexable text."""


def load_document(filename: str, content: bytes) -> list[DocumentPage]:
    """Extract text from a PDF or DOCX byte stream; PDF pages remain separate."""
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        try:
            with pymupdf.open(stream=content, filetype="pdf") as pdf:
                pages = [DocumentPage(clean_text(page.get_text("text")), index + 1) for index, page in enumerate(pdf)]
        except Exception as error:
            raise DocumentLoadError("This PDF could not be opened. Check that it is a valid, unencrypted PDF.") from error
        pages = [page for page in pages if page.text]
    elif extension == ".docx":
        try:
            document = Document(BytesIO(content))
            sections = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if values:
                        sections.append(" | ".join(values))
            text = clean_text("\n".join(sections))
            pages = [DocumentPage(text)] if text else []
        except Exception as error:
            raise DocumentLoadError("This DOCX could not be opened. Check that it is a valid Word document.") from error
    else:
        raise DocumentLoadError("Unsupported file type. Upload a PDF or DOCX document.")

    if not pages:
        raise DocumentLoadError("No selectable text was found. This may be an image-only PDF; OCR is not included.")
    return pages
